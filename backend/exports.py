"""Excel / CSV writers for reports. Pure functions - no database, no FastAPI.

A report is described by ``columns`` = [(key, label, kind)] where kind is one of
"text", "money", "number", "date", plus ``rows`` (list of dicts) and an optional
``summary`` dict (totals) written as a bold last row.
"""
import csv
from datetime import date, datetime, time
from decimal import Decimal
from io import BytesIO, StringIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
CSV_MIME = "text/csv; charset=utf-8"
_FORMULA_STARTS = ("=", "+", "-", "@", "\t", "\r")


def safe_text(value):
    """Stop spreadsheet formula injection: text that starts with = + - @ is prefixed
    with an apostrophe so Excel shows it instead of executing it."""
    if isinstance(value, str) and value.startswith(_FORMULA_STARTS):
        return "'" + value
    return value


def _cell(value, kind):
    if value is None:
        return None
    if kind in ("money", "number") and isinstance(value, (Decimal, int, float)):
        return float(value)
    if isinstance(value, (datetime, date, time)):
        return value.isoformat() if isinstance(value, time) else value
    if isinstance(value, Decimal):
        return float(value)
    return safe_text(value)


def build_xlsx(title: str, columns: list, rows: list, summary: dict | None = None, subtitle: str = "") -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    ws.append([title])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([subtitle])
    ws["A2"].font = Font(italic=True, color="64748B")
    ws.append([])
    header_row = ws.max_row + 1
    ws.append([label for _, label, _ in columns])
    for c in range(1, len(columns) + 1):
        cell = ws.cell(row=header_row, column=c)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2563EB")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row in rows:
        ws.append([_cell(row.get(key), kind) for key, _, kind in columns])
    if summary is not None:
        ws.append([_cell(summary.get(key), kind) if key in summary else None for key, _, kind in columns])
        for c in range(1, len(columns) + 1):
            ws.cell(row=ws.max_row, column=c).font = Font(bold=True)

    for idx, (_, _, kind) in enumerate(columns, start=1):
        for r in range(header_row + 1, ws.max_row + 1):
            cell = ws.cell(row=r, column=idx)
            if kind == "money":
                cell.number_format = "#,##0.00"
            elif kind == "number":
                cell.number_format = "0.##"
            elif kind == "date":
                cell.number_format = "yyyy-mm-dd"
        width = max([len(str(columns[idx - 1][1]))] +
                    [len(str(ws.cell(row=r, column=idx).value or "")) for r in range(header_row + 1, ws.max_row + 1)])
        ws.column_dimensions[get_column_letter(idx)].width = min(max(width + 2, 10), 40)
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_csv(columns: list, rows: list, summary: dict | None = None) -> bytes:
    out = StringIO()
    writer = csv.writer(out)
    writer.writerow([label for _, label, _ in columns])
    for row in rows:
        writer.writerow([_csv_value(row.get(key)) for key, _, _ in columns])
    if summary is not None:
        writer.writerow([_csv_value(summary.get(key)) if key in summary else "" for key, _, _ in columns])
    return ("\ufeff" + out.getvalue()).encode("utf-8")      # BOM so Excel reads UTF-8 correctly


def _csv_value(value):
    if value is None:
        return ""
    if isinstance(value, Decimal):
        return f"{value:.2f}"
    return safe_text(value)